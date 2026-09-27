from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from src.exporters import export_csv, export_excel


def test_csv_has_no_index():
    data = pd.DataFrame({"Player": ["Alpha"], "PTS": [12.5]})
    text = export_csv(data).decode("utf-8-sig")
    assert text == "Player,PTS\nAlpha,12.5\n"


def test_excel_formatting():
    data = pd.DataFrame({"Player": ["Alpha"], "FG%": [0.5]})
    workbook = load_workbook(BytesIO(export_excel(data)))
    sheet = workbook["NBA Stats"]
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:B2"
    assert sheet["A1"].font.bold is True
    assert sheet["B2"].number_format == "0.0%"


def test_excel_supports_dataset_specific_sheet_name():
    workbook = load_workbook(BytesIO(export_excel(pd.DataFrame({"Team": ["DEN"]}), "Standings")))
    assert workbook.sheetnames == ["Standings"]


def test_international_player_names_survive_csv_and_excel_exports():
    data = pd.DataFrame({"Player": ["Luka Dončić", "Nikola Jokić"]})

    csv_text = export_csv(data).decode("utf-8-sig")
    workbook = load_workbook(BytesIO(export_excel(data)))
    sheet = workbook["NBA Stats"]

    assert "Luka Dončić" in csv_text
    assert "Nikola Jokić" in csv_text
    assert sheet["A2"].value == "Luka Dončić"
    assert sheet["A3"].value == "Nikola Jokić"
