"""CSV and formatted Excel export functions."""

from __future__ import annotations

from io import BytesIO

import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def export_csv(dataframe: pd.DataFrame) -> bytes:
    """Create UTF-8 CSV bytes without a DataFrame index."""
    return dataframe.to_csv(index=False).encode("utf-8-sig")


def export_excel(dataframe: pd.DataFrame, sheet_name: str = "NBA Stats") -> bytes:
    """Create a compact, readable Excel workbook in memory."""
    safe_sheet_name = sheet_name[:31] or "NBA Stats"
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        dataframe.to_excel(writer, sheet_name=safe_sheet_name, index=False)
        worksheet = writer.sheets[safe_sheet_name]
        worksheet.freeze_panes = "A2"

        for cell in worksheet[1]:
            cell.font = Font(bold=True)

        if worksheet.max_column:
            worksheet.auto_filter.ref = worksheet.dimensions

        for index, column in enumerate(dataframe.columns, start=1):
            values = dataframe[column].dropna().astype(str).tolist()
            longest = max([len(str(column)), *(len(value) for value in values)], default=len(str(column)))
            worksheet.column_dimensions[get_column_letter(index)].width = min(longest + 2, 40)

            if str(column).endswith("%"):
                for row in range(2, worksheet.max_row + 1):
                    worksheet.cell(row=row, column=index).number_format = "0.0%"

    return output.getvalue()
